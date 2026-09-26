"""Build the client proposal from its Markdown source using bundled python-docx."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs' / 'client-proposal.md'
OUTPUT = ROOT / 'deliverables' / 'LK_Group_Client_Implementation_Plan_KO.docx'
FONT = 'Malgun Gothic'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.8)
section.left_margin, section.right_margin = Cm(2.0), Cm(2.0)
section.footer_distance = Cm(0.8)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet']:
    style = doc.styles[name]
    style.font.name = FONT
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), FONT)
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'en-AU'); lang.set(qn('w:eastAsia'), 'ko-KR')
    style._element.rPr.append(lang)

normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.2
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True
for name, size, before, after in [('Title', 26, 0, 14), ('Heading 1', 17, 0, 12), ('Heading 2', 12, 9, 6)]:
    style = doc.styles[name]
    style.font.size = Pt(size)
    style.font.bold = name != 'Title'
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.line_spacing = 1.15
    ppr = style._element.get_or_add_pPr()
    for x in list(ppr):
        if x.tag in [qn('w:pBdr'), qn('w:shd')]: ppr.remove(x)

def font_run(run, size=None, bold=None, color='000000'):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    run.font.color.rgb = RGBColor.from_string(color)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold

def link(p, label, url):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    f = OxmlElement('w:rFonts')
    for key in ['ascii', 'hAnsi', 'eastAsia']: f.set(qn('w:'+key), FONT)
    pr.append(f)
    c = OxmlElement('w:color'); c.set(qn('w:val'), '254C70'); pr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); pr.append(u)
    r.append(pr); t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)

def rich(p, text, size=None, color='000000'):
    pattern = r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*.+?\*\*)'
    for part in re.split(pattern, text):
        if not part: continue
        m = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', part)
        if m:
            link(p, m.group(1), m.group(2))
        else:
            bold = part.startswith('**') and part.endswith('**')
            r = p.add_run(part[2:-2] if bold else part)
            font_run(r, size, True if bold else None, color)

def widths(rows):
    n = len(rows[0])
    if n == 2: return [4.5, 12.5]
    if n == 3:
        if rows[0][0] in ['범위', '단계']: return [4.6, 4.0, 8.4]
        return [3.1, 7.6, 6.3]
    if n == 4:
        if rows[0][0] == '단계': return [3.2, 2.8, 8.0, 3.0]
        return [2.9, 4.0, 6.2, 3.9]
    return [17/n]*n

def add_table(rows):
    w = widths(rows)
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, value in zip(table.columns, w): col.width = Cm(value)
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: el.set(qn('w:'+k), v)
        borders.append(el)
    props.append(borders)
    for i, values in enumerate(rows):
        row = table.add_row()
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for j, text in enumerate(values):
            cell = row.cells[j]; cell.width = Cm(w[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, val in [('top','95'),('bottom','95'),('left','105'),('right','105')]:
                e = OxmlElement('w:'+side); e.set(qn('w:w'),val); e.set(qn('w:type'),'dxa'); margins.append(e)
            tcpr.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '253B50' if i == 0 else ('F1F4F6' if i%2 == 0 else 'FFFFFF')); tcpr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.14
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.keep_with_next = i == 0
            if i == 0 or (len(text) < 19 and j == 0): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rich(p, text, size=9.5, color='FFFFFF' if i == 0 else '000000')
            if i == 0:
                for r in p.runs: r.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 0.3
    font_run(p.add_run(''), 2)

lines = SOURCE.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    s = lines[i].strip()
    if not s: i+=1; continue
    if s == '<!-- pagebreak -->':
        doc.add_page_break(); i+=1; continue
    if s.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{3,}:?',c) for c in cells): rows.append(cells)
            i+=1
        add_table(rows); continue
    if s.startswith('# '):
        p = doc.add_paragraph(style='Title'); rich(p, s[2:])
    elif s.startswith('### '):
        p = doc.add_paragraph(style='Heading 2'); rich(p, s[4:])
    elif s.startswith('## '):
        p = doc.add_paragraph(style='Heading 1'); rich(p, s[3:])
    elif s.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); rich(p,s[2:])
        p.paragraph_format.space_after = Pt(4)
    else:
        p=doc.add_paragraph(); rich(p,s)
        if s == 'L&K Group':
            for r in p.runs: font_run(r, 16, True)
        elif s.startswith('클라이언트 제안서'):
            for r in p.runs: font_run(r, 9, False, '555555')
            p.paragraph_format.space_after = Pt(16)
    i+=1

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font_run(footer.add_run('L&K Group  |  '), 8, False, '666666')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = '웹 서비스 구축 및 마케팅 실행 계획'
doc.core_properties.subject = 'L&K Group 통합 웹 서비스와 쇼핑몰 및 운영관리 구축'
doc.core_properties.author = 'L&K Group'
doc.core_properties.keywords = '웹사이트, 견적, 예약, 쇼핑몰, ERP, CRM, 마케팅'
OUTPUT.parent.mkdir(exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT.name)
print(f'Paragraphs={len(doc.paragraphs)} Tables={len(doc.tables)} Planned pages={1+SOURCE.read_text(encoding="utf-8").count("<!-- pagebreak -->")}')
